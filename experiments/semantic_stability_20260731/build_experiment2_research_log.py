from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from build_experiment1_research_log import _add_body, _add_figure, _add_heading, _add_table, _add_title


def build_log(template_path: Path, output_path: Path, result_dir: Path) -> None:
    summary = json.loads((result_dir / "semantic_candidate_rerank_summary.json").read_text(encoding="utf-8"))
    report_ccw = summary["per_sequence"]["Aisle_CCW"]
    report_cw = summary["per_sequence"]["Aisle_CW"]
    combined = summary["combined"]
    availability = summary["candidate_availability"]
    document = Document(template_path)
    from build_experiment1_research_log import _clear_document_body

    _clear_document_body(document)
    _add_title(document, "研究日志（2026.07.31）")

    _add_heading(document, "一、实验目的")
    _add_body(
        document,
        "实验一证明了跨日期语义类别具有长期一致性，但机器人当前位置附近的总体稳定性不能直接解释定位误差。原因是重复货架区域本身也很稳定。实验二据此把语义信息从全局置信度改为候选级证据：对于同一帧粗定位提供的 TopK 候选，分别判断当前点云语义锚点在不同候选位姿下是否能与历史语义地图一致，再验证该信息能否帮助选对候选 patch。"
    )

    _add_heading(document, "二、实验设置")
    _add_body(
        document,
        "本实验固定 Oct.12 Aisle_CCW 和 Aisle_CW 的既有 v14a 候选结果：粗定位 TopK、深度匹配分数、三种初始化和每个候选的 ICP 位姿均不重新计算。对每个候选，仅将当前帧 LiDAR 点投影到左右语义分割结果，再按该候选的最终位姿投到实验一构建的 Jun.15+Jun.23 细粒度语义地图中。只有当前标签与栅格历史标签分布一致、且落在高可靠区域的点才提供较高语义分数；半动态类低权重参与，动态类不参与。"
    )
    _add_table(
        document,
        ["条件", "候选选择方式", "作用"],
        [
            ["原候选分数", "沿用冻结的候选 final score", "候选级基线"],
            ["语义融合", "final score + 0.20 × 语义一致性分数", "预先固定的轻量融合，不在 Oct.12 调参"],
            ["仅语义", "只按语义一致性分数排序", "检验语义能否替代几何/深度的失败对照"],
            ["TopK 上限", "事后选择候选中误差最小者", "仅用于衡量候选集合中的可达上限"],
        ],
        [1500, 4000, 2795],
    )

    _add_heading(document, "三、候选级结果")
    _add_body(
        document,
        f"在冻结候选上，Aisle_CCW 的原候选分数与语义融合平均误差均为 {report_ccw['candidate_final_score']['mean_position_error_m']:.6f} m，说明该方向没有发生有益排序变化；Aisle_CW 从 {report_cw['candidate_final_score']['mean_position_error_m']:.6f} m 降至 {report_cw['semantic_blend']['mean_position_error_m']:.6f} m，改善 {report_cw['candidate_final_score']['mean_position_error_m'] - report_cw['semantic_blend']['mean_position_error_m']:.6f} m。合并 2007 帧后，平均误差从 {combined['candidate_final_score']['mean_position_error_m']:.6f} m 变为 {combined['semantic_blend']['mean_position_error_m']:.6f} m，改善仅 {combined['candidate_final_score']['mean_position_error_m'] - combined['semantic_blend']['mean_position_error_m']:.6f} m，且小于 0.5 m 的比例没有变化。"
    )
    _add_body(
        document,
        f"候选集合本身的可用性很高：CCW 的候选位姿成功 Recall@3 为 {availability['Aisle_CCW']['recall_at_3_pose_success'] * 100:.2f}%、MRR 为 {availability['Aisle_CCW']['mrr_pose_success']:.4f}；CW 分别为 {availability['Aisle_CW']['recall_at_3_pose_success'] * 100:.2f}% 和 {availability['Aisle_CW']['mrr_pose_success']:.4f}。同时，TopK 事后上限的合并平均误差为 {combined['oracle_topk']['mean_position_error_m']:.6f} m，明显低于当前原候选分数，说明候选中仍有可利用的正确位姿，但本版静态语义一致性尚未充分识别出它们。"
    )
    _add_figure(document, result_dir / "semantic_candidate_rerank_comparison.png", "图1  冻结 TopK 下的候选级语义重排结果；右侧单独显示仅语义排序的失败尺度")

    _add_heading(document, "四、失败分析")
    _add_body(
        document,
        f"仅按语义分数选择候选时，CCW 和 CW 的平均误差分别达到 {report_ccw['semantic_only']['mean_position_error_m']:.3f} m 与 {report_cw['semantic_only']['mean_position_error_m']:.3f} m，并伴随大量 patch 切换。这说明仓库中“看到货架、地面、墙柱”等语义结构并不足以确定绝对位置：相邻货架区通常具有相似的语义组成，语义没有几何形状、相对位姿和时序约束便会产生新的歧义。"
    )
    _add_body(
        document,
        "语义融合的提升极小也具有研究价值。当前候选 final score 已融合深度、BEV、ICP 与时序，且大多数帧的正确候选本来就在第一名；语义分数只在少数近分候选中触发重排。下一版不应继续提高静态语义权重，因为这会向仅语义失败对照靠近，而应学习“什么时候语义能区分候选、什么时候应保持原始几何排序”的可靠性门控。"
    )

    _add_heading(document, "五、下一步")
    _add_body(
        document,
        "后续实验将不把语义分数直接相加，而是构造任务对齐的学习式可靠性门控。它的输入同时包括深度候选间隔、深度初始化与 tracker 初始化的 ICP 残差差、当前语义锚点覆盖率、历史稳定性和最近多帧的位姿创新量，输出采用深度初始化、tracker 初始化或保持状态的连续权重。训练标签由训练日期上的反事实配准结果生成，而不是由单一人员面积阈值生成。"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the Experiment 2 research log from the retained July 18 template.")
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--result-dir", required=True)
    args = parser.parse_args()
    build_log(Path(args.template), Path(args.output), Path(args.result_dir))


if __name__ == "__main__":
    main()
