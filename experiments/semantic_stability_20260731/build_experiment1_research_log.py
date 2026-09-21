from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BODY_FONT = "宋体"
TITLE_FONT = "黑体"
TABLE_WIDTH_DXA = 8295


def _set_font(run, name: str, size_pt: float, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    _set_font(paragraph.add_run(text), TITLE_FONT, 14, bold=True)


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    _set_font(paragraph.add_run(text), BODY_FONT, 12, bold=True)


def _add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    _set_font(paragraph.add_run(text), BODY_FONT, 12)


def _set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(width_dxa))
    tc_width.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pr = table._tbl.tblPr
    table_width = table_pr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths_dxa[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_cell_text(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.2
    paragraph.clear()
    _set_font(paragraph.add_run(text), BODY_FONT, 10.5, bold=bold)


def _add_table(document: Document, headers: list[str], rows: list[list[str]], widths_dxa: list[int]) -> None:
    table = document.add_table(rows=1, cols=len(headers), style="Table Grid")
    _set_table_geometry(table, widths_dxa)
    for index, text in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], text, bold=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E7E6E6")
        table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shading)
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            _set_cell_text(cells[index], text)
    _set_table_geometry(table, widths_dxa)


def _add_figure(document: Document, image_path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(image_path), width=Inches(5.72))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(3)
    caption_paragraph.paragraph_format.space_after = Pt(0)
    caption_paragraph.paragraph_format.line_spacing = 1.0
    caption_run = caption_paragraph.add_run(caption)
    _set_font(caption_run, BODY_FONT, 10.5)
    caption_run.font.color.rgb = RGBColor(89, 89, 89)


def _format_ratio(value: float) -> str:
    return f"{value * 100:.2f}%"


def build_log(template_path: Path, output_path: Path, result_dir: Path) -> None:
    summary = json.loads((result_dir / "semantic_stability_summary.json").read_text(encoding="utf-8"))
    correlation = json.loads((result_dir / "stability_localization_correlation.json").read_text(encoding="utf-8"))
    document = Document(template_path)
    _clear_document_body(document)
    _add_title(document, "研究日志（2026.07.31）")

    _add_heading(document, "一、实验目的")
    _add_body(
        document,
        "本次实验围绕仓库场景的中长期变化开展。前期结果表明，行人等短期遮挡只能解释部分定位误差，而重复货架、货物堆放和临时设备的长期变化同样会影响候选 patch 的可靠性。因此，本实验不直接修改原始点云地图，而是在原有几何地图外构建一层可持续累积的语义稳定性地图，用于记录不同区域在重复观测中更接近稳定结构还是半动态/动态结构。"
    )
    _add_body(
        document,
        "实验一首先回答两个问题：不同日期的 Aisle 观测能否形成一致的稳定性统计；以及这种区域级稳定性是否可以直接解释已有定位器的误差。第二个问题十分重要，因为如果总体稳定性已经足够预测误差，后续只需增加全局权重；反之，则必须把语义信息放到候选 patch 的匹配一致性中使用。"
    )

    _add_heading(document, "二、数据划分与方法")
    _add_body(
        document,
        "初始语义地图只使用 Jun.15 的 Aisle_CCW_Run_1 和 Aisle_CW_Run_1；同日验证使用 Jun.15 的两个 Run_2；跨日期更新使用 Jun.23 的两个 Run_1。每隔 5 帧采样一次点云，将裁剪后的 LiDAR 点投影到左右语义分割图，再利用对应日期的位姿投到统一地图坐标系。Oct.12 数据不参与建图和阈值选择，只在地图完成后用于离线诊断。"
    )
    _add_body(
        document,
        "每个 0.5 m 栅格以帧为单位记录主导类别，避免单帧中高密度点云压倒重复观测。墙、柱、货架、地面、天花板和固定机械归为高可信静态类；货物、推车、路锥和文本区域归为中可信半动态类；人员、叉车及其他动态物体归为低可信动态类。稳定性同时考虑静态证据比例和该栅格的重复观测次数。"
    )
    _add_table(
        document,
        ["数据角色", "序列与用途", "处理方式"],
        [
            ["初始建图", "Jun.15 Aisle_CCW_Run_1、Aisle_CW_Run_1", "构建初始稳定性层，不使用测试数据"],
            ["同日验证", "Jun.15 Aisle_CCW_Run_2、Aisle_CW_Run_2", "检查重复观测的一致性，不写入初始地图"],
            ["跨日期更新", "Jun.23 Aisle_CCW_Run_1、Aisle_CW_Run_1", "模拟一次历史地图更新"],
            ["最终诊断", "Oct.12 Aisle_CCW、Aisle_CW", "仅分析已冻结基线结果，不参与建图或调参"],
        ],
        [1700, 3800, 2795],
    )

    _add_heading(document, "三、跨日期语义稳定性统计")
    initial = summary["role_reports"]["initial"]
    update = summary["role_reports"]["cross_day_update"]
    initial_total = sum(initial["cell_group_evidence"].values())
    update_total = sum(update["cell_group_evidence"].values())
    same_day = summary["consistency"]["same_day"]
    cross_day = summary["consistency"]["cross_day"]
    _add_body(
        document,
        f"初始地图覆盖 {summary['initial_map']['populated_cells']} 个语义栅格；合入 Jun.23 高置信历史观测后覆盖 {summary['updated_map']['populated_cells']} 个栅格，新增 {summary['updated_map']['populated_cells'] - summary['initial_map']['populated_cells']} 个栅格。初始建图的帧级主导证据中，静态类占 {_format_ratio(initial['cell_group_evidence']['static'] / initial_total)}，半动态类占 {_format_ratio(initial['cell_group_evidence']['semi_dynamic'] / initial_total)}，动态类仅占 {_format_ratio(initial['cell_group_evidence']['dynamic'] / initial_total)}；Jun.23 跨日期观测中静态证据仍占 {_format_ratio(update['cell_group_evidence']['static'] / update_total)}。"
    )
    _add_body(
        document,
        f"初始 Run_1 与同日 Run_2 在 {same_day['eligible_cells']} 个有效共享栅格上的主导类别一致率为 {_format_ratio(same_day['dominant_group_agreement'])}，跨日期的 Jun.15 Run_1 与 Jun.23 Run_1 在 {cross_day['eligible_cells']} 个有效共享栅格上的一致率仍为 {_format_ratio(cross_day['dominant_group_agreement'])}。跨日期相对同日只下降 {(same_day['dominant_group_agreement'] - cross_day['dominant_group_agreement']) * 100:.2f} 个百分点，说明稳定结构的长期语义信号确实存在。"
    )
    _add_figure(document, result_dir / "semantic_stability_initial_vs_updated.png", "图1  Jun.15 初始语义稳定性地图与合入 Jun.23 观测后的更新结果")
    _add_figure(document, result_dir / "semantic_stability_evidence.png", "图2  三个时间阶段的帧级主导语义证据比例与重复观测一致性")

    _add_heading(document, "四、与现有定位误差的关系")
    _add_body(
        document,
        "为避免把语义稳定性直接调到 Oct.12 测试集上，本节只做事后诊断：对每个测试帧，以真值位置周围 7.5 m 内的 Jun.15+Jun.23 稳定性栅格计算局部平均稳定性，并与已经完成的 v14a 基线误差比较。该稳定性不作为定位输入，也不参与任何超参数选择。"
    )
    overall_r = correlation["overall_pearson_stability_error"]
    ccw = correlation["per_sequence"]["Aisle_CCW"]
    cw = correlation["per_sequence"]["Aisle_CW"]
    _add_body(
        document,
        f"结果显示总体 Pearson 相关系数仅为 {overall_r:.4f}，其中 CCW 为 {ccw['pearson_stability_error']:.4f}，CW 为 {cw['pearson_stability_error']:.4f}。这不是预期中的强负相关。特别是 CCW 的高稳定性四分位中仍包含 {ccw['high_stability_quartile']['position_error_above_0p5m']} 个大于等于 0.5 m 的误差帧，CW 中也包含 {cw['high_stability_quartile']['position_error_above_0p5m']} 个。可见货架等结构在长期语义上稳定，并不代表它们在几何上没有重复歧义。"
    )
    _add_figure(document, result_dir / "stability_error_scatter.png", "图3  历史语义稳定性与 Oct.12 基线定位误差的事后诊断关系")

    _add_heading(document, "五、实验结论与下一步")
    _add_body(
        document,
        "实验一得到的正面结果是：跨日期地图中高可信静态结构占主导，且主导语义类别在同日和跨日均保持约 98% 的一致性，因此用语义稳定性描述地图区域的长期可靠性是可行的。与此同时，实验一也得到一个必须保留的负面结论：机器人当前位置附近的总体稳定性不能直接作为最终定位置信度，因为重复货架区本身很稳定，却仍可能使多个候选 patch 同时获得较高几何支持。"
    )
    _add_body(
        document,
        "因此，下一步不应把稳定性简单加到全局最终分数，而应开展候选级语义重排实验：固定粗定位 TopK 和 ICP，分别计算当前帧稳定语义锚点与每个候选 patch 的长期语义分布是否一致，再观察正确 patch 的 Top1、MRR、错误 patch 切换率和完整 Aisle 精度是否改善。这样才能验证语义稳定性是否真正帮助系统在重复结构中选对位置。"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the Experiment 1 research log from the retained July 18 template.")
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--result-dir", required=True)
    args = parser.parse_args()
    build_log(Path(args.template), Path(args.output), Path(args.result_dir))


if __name__ == "__main__":
    main()
