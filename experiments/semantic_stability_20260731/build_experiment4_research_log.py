from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont

from experiments.semantic_stability_20260731.build_experiment1_research_log import (
    _add_body,
    _add_figure,
    _add_heading,
    _add_table,
    _add_title,
    _clear_document_body,
)


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def _draw_summary(output_path: Path, validation: dict, blind: dict) -> None:
    canvas = Image.new("RGB", (1420, 910), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = _font(31, True)
    body_font = _font(22)
    small_font = _font(18)
    draw.text((48, 30), "实验四：多帧静态点融合 + 截尾 ICP", font=title_font, fill="#1d2939")
    draw.text((48, 79), "固定参数先在 Jun.23 验证，再在 Oct.12 盲测；两阶段均按每 8 帧重放", font=body_font, fill="#475467")
    panels = [(65, 155, "跨日期验证（Jun.23）", validation["combined"]), (745, 155, "盲测（Oct.12）", blind["combined"])]
    for left, top, title, result in panels:
        width, height = 595, 320
        draw.rounded_rectangle((left, top, left + width, top + height), radius=15, outline="#cbd5e1", width=2)
        draw.text((left + 25, top + 22), title, font=body_font, fill="#1d2939")
        entries = [
            ("原单帧", result["baseline"]["mean_position_error_m"], "#64748b"),
            ("多帧静态", result["multiframe_static"]["mean_position_error_m"], "#2563eb"),
        ]
        max_value = max(value for _, value, _ in entries) * 1.25
        for index, (label, value, color) in enumerate(entries):
            x = left + 120 + index * 210
            bar_height = 165 * value / max_value
            draw.rounded_rectangle((x, top + 245 - bar_height, x + 100, top + 245), radius=6, fill=color)
            draw.text((x - 4, top + 214 - bar_height), f"{value:.4f} m", font=small_font, fill="#344054")
            draw.text((x + 2, top + 265), label, font=small_font, fill="#344054")
        draw.text(
            (left + 24, top + 286),
            f"平均改善 {result['mean_improvement_m']:+.4f} m；95 分位 {result['baseline']['p95_position_error_m']:.4f} → {result['multiframe_static']['p95_position_error_m']:.4f} m",
            font=small_font,
            fill="#475467",
        )
    draw.rounded_rectangle((65, 535, 1355, 825), radius=15, outline="#cbd5e1", width=2)
    draw.text((92, 560), "Oct.12 分路线结果", font=body_font, fill="#1d2939")
    headers = ["路线", "原平均误差", "多帧误差", "改善", "接受率", "改善帧比例"]
    x_positions = [92, 280, 485, 682, 856, 1040]
    for x, header in zip(x_positions, headers, strict=True):
        draw.text((x, 614), header, font=small_font, fill="#667085")
    for row_index, sequence_name in enumerate(("Aisle_CCW", "Aisle_CW")):
        result = blind["per_sequence"][sequence_name]
        values = [
            sequence_name,
            f"{result['baseline']['mean_position_error_m']:.4f} m",
            f"{result['multiframe_static']['mean_position_error_m']:.4f} m",
            f"{result['mean_improvement_m']:+.4f} m",
            f"{result['accepted_rate'] * 100:.2f}%",
            f"{result['improved_frame_ratio'] * 100:.2f}%",
        ]
        y = 663 + row_index * 62
        draw.line((88, y - 14, 1328, y - 14), fill="#e2e8f0", width=1)
        for x, value in zip(x_positions, values, strict=True):
            draw.text((x, y), value, font=body_font, fill="#344054")
    draw.text((92, 775), "使用过去 2 个已定位帧与当前帧的语义动态点剔除后点云；不读取未来帧或真值。", font=small_font, fill="#475467")
    canvas.save(output_path)


def build_log(template_path: Path, output_path: Path, result_dir: Path) -> None:
    validation = json.loads((result_dir / "multiframe_static_validation_summary.json").read_text(encoding="utf-8"))
    blind = json.loads((result_dir / "multiframe_static_oct12_blind_summary.json").read_text(encoding="utf-8"))
    figure_path = result_dir / "multiframe_static_summary.png"
    _draw_summary(figure_path, validation, blind)
    document = Document(template_path)
    _clear_document_body(document)
    _add_title(document, "研究日志（2026.07.31）")

    _add_heading(document, "一、实验目的")
    _add_body(
        document,
        "实验三说明，仅凭单帧的深度置信度、ICP 残差和语义面积，无法可靠预测错误切换。进一步观察定位尖峰可知，人员或推车经过时，单帧点云中有效货架边缘会被部分遮挡；即使动态点已经剔除，余下的静态点数量和视角也不足以稳定约束 ICP。本实验不再学习“选哪条初始化”，而是直接提升同一候选位姿周围的几何可观测性：把当前帧与过去两个已定位帧的静态点投到世界坐标系中融合，再用截尾 ICP 估计一个小的全局校正量。")

    _add_heading(document, "二、固定方法与防泄漏设置")
    _add_body(
        document,
        "每个采样时刻只使用当前帧及其过去两个采样时刻的点云；每帧先通过双目语义分割去除人员、叉车和其他动态类别的点，再按已有 v14a 在线位姿变换到世界坐标系。融合点云与六月十五日建立的全局地图进行 20 次截尾 ICP：在 0.8 m 对应距离内仅保留残差最小的 70% 对应关系。为了防止错误轨迹被放大，只有内点比例不少于 0.55、平移校正不超过 0.45 m、偏航校正不超过 3° 时才接纳校正，否则保留原位姿。所有这些参数在 Oct.12 之前固定。")
    _add_table(
        document,
        ["环节", "输入", "处理", "安全约束"],
        [
            ["静态点选择", "当前及过去两帧 LiDAR、双目语义分割", "去除类别 12–15 的动态点", "只读取当前或过去帧"],
            ["短时融合", "三帧静态点与已有在线位姿", "统一到世界坐标后体素化", "历史位姿距当前超过 1.5 m 不参与"],
            ["鲁棒细配准", "融合点与全局点云地图", "截尾 ICP，仅保留最小残差的 70%", "校正幅度、内点比例和有限残差三重门限"],
        ],
        [1450, 2450, 2600, 2745],
    )

    _add_heading(document, "三、跨日期验证")
    validation_combined = validation["combined"]
    _add_body(
        document,
        f"在 Jun.23 的 278 个跨日期采样帧上，原轨迹的平均误差为 {validation_combined['baseline']['mean_position_error_m']:.6f} m，多帧静态细配准后为 {validation_combined['multiframe_static']['mean_position_error_m']:.6f} m，改善 {validation_combined['mean_improvement_m']:.6f} m；95 分位误差由 {validation_combined['baseline']['p95_position_error_m']:.6f} m 降至 {validation_combined['multiframe_static']['p95_position_error_m']:.6f} m。CCW 与 CW 两条路线均有改善，接纳率为 {validation_combined['accepted_rate'] * 100:.2f}%，且 {validation_combined['improved_frame_ratio'] * 100:.2f}% 的帧误差下降。因此该方法通过跨日期准入，可以按不变参数进入 Oct.12 盲测。")

    _add_heading(document, "四、Oct.12 盲测结果")
    blind_combined = blind["combined"]
    _add_body(
        document,
        f"在不改动任何参数的 Oct.12 Aisle 盲测中，合并的平均误差从 {blind_combined['baseline']['mean_position_error_m']:.6f} m 降至 {blind_combined['multiframe_static']['mean_position_error_m']:.6f} m，改善 {blind_combined['mean_improvement_m']:.6f} m；中位数从 {blind_combined['baseline']['median_position_error_m']:.6f} m 降至 {blind_combined['multiframe_static']['median_position_error_m']:.6f} m，95 分位从 {blind_combined['baseline']['p95_position_error_m']:.6f} m 降至 {blind_combined['multiframe_static']['p95_position_error_m']:.6f} m。小于 0.5 m 的比例保持 {blind_combined['baseline']['below_0p5m'] * 100:.4f}%，说明收益主要来自压低常规误差和误差尖峰，而不是改变已经很高的粗阈值成功率。")
    _add_figure(document, figure_path, "图 1  多帧静态点融合与截尾 ICP 的跨日期验证、Oct.12 盲测及分路线结果")

    _add_heading(document, "五、结论与边界")
    _add_body(
        document,
        "本实验说明，面对人员造成的局部遮挡，与其用单帧分类器猜测“应不应该信任 tracker”，不如直接增加可用于配准的静态几何证据，并在对应关系层面抑制异常残差。该机制在验证和盲测上方向一致，是目前语义辅助精定位中第一条通过跨日期准入的优化。不过，本结果仍是基于既有 v14a 在线轨迹的因果重放：多帧融合使用的是过去已经输出的位姿，尚未替换主线实时循环。因此下一步需把短时静态子图缓存接入在线局部化器，验证每帧顺序重放时能否保持相同收益，并进一步研究长期地图更新时如何只吸收被多次观测证实的稳定语义。")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the Experiment 4 research log from the July 18 template.")
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--result-dir", required=True)
    args = parser.parse_args()
    build_log(Path(args.template), Path(args.output), Path(args.result_dir))


if __name__ == "__main__":
    main()
