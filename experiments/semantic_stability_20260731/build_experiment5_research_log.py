from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from experiments.semantic_stability_20260731.build_experiment1_research_log import (
    _add_body,
    _add_figure,
    _add_heading,
    _add_table,
    _add_title,
    _clear_document_body,
)


def build_log(template_path: Path, output_path: Path, result_dir: Path) -> None:
    summary = json.loads((result_dir / "semantic_update_safety_summary.json").read_text(encoding="utf-8"))
    policies = summary["policies"]
    strict = policies["persistent_label"]
    document = Document(template_path)
    _clear_document_body(document)
    _add_title(document, "研究日志（2026.07.31）")

    _add_heading(document, "一、实验目的")
    _add_body(
        document,
        "长期地图更新的风险不在于是否能够增加更多点，而在于一次错误写入会把临时货物、推车或遮挡痕迹固化为地图结构，随后持续影响粗定位候选和精定位对应关系。实验一证明静态语义在跨日期上总体稳定，但并不意味着每一次当前观测都应写入。因此本实验把“更新”限定为向语义可靠性层增加一条候选证据，而不修改原始几何点云地图；目标是比较不同安全门控在覆盖率与未来可验证性之间的权衡。")

    _add_heading(document, "二、评估协议")
    _add_body(
        document,
        "决策侧只读取 Jun.15+Jun.23 Run1 已建立的历史语义稳定性层，以及 Jun.23 Run2 的当前观测。每 0.5 m 单元由当前点云投影到双目语义分割后得到细类别和跨帧支持次数。Oct.12 Aisle 完全不参与任何门限选择，仅用于未来确认：若一个被接受的单元在 Oct.12 中仍可观测，且其主导细类别仍与待写入观测一致、并属于静态类别，则判为未来确认安全；否则判为矛盾。未在 Oct.12 覆盖到的单元只记为不可确认，不计入确认率分母。")
    _add_table(
        document,
        ["策略", "接受条件", "作用"],
        [
            ["无筛选", "任何有语义的当前单元", "覆盖最高，但混入动态和类别漂移"],
            ["仅当前静态", "当前主导类别属于静态类", "排除显式动态与半动态观测"],
            ["静态+历史稳定", "当前静态，历史主导组静态，稳定度≥0.70，历史证据≥3", "要求当前与长期地图均支持静态"],
            ["稳定+跨帧同类", "前一策略加当前至少 3 次观测、与历史细类别一致", "只写入重复验证过的细粒度结构"],
        ],
        [1750, 3950, 2595],
    )

    _add_heading(document, "三、安全性结果")
    _add_body(
        document,
        f"Jun.23 Run2 共产生 {summary['num_current_cells']} 个待评估单元，Oct.12 中有 {summary['num_future_cells']} 个单元被重新观测。无筛选策略接受全部单元，未来可观测单元的细类别确认率只有 {policies['naive']['future_confirmation_rate'] * 100:.2f}%；只检查当前静态类别后提升到 {policies['semantic_static']['future_confirmation_rate'] * 100:.2f}%。加入历史稳定度与历史静态主导类别后，确认率达到 {policies['stable_static']['future_confirmation_rate'] * 100:.2f}%，说明长期证据确实能够过滤部分当前偶发观测。")
    _add_body(
        document,
        f"最严格的“稳定+跨帧同类”策略接受 {strict['accepted_cells']} 个单元，占全部候选的 {strict['coverage_over_current_cells'] * 100:.2f}%；其中 {strict['future_observed_cells']} 个在 Oct.12 中可复查，细类别确认率为 {strict['future_confirmation_rate'] * 100:.2f}%，未来矛盾率降至 {strict['future_contradiction_rate'] * 100:.2f}%。相较无筛选，确认率提高 {(strict['future_confirmation_rate'] - policies['naive']['future_confirmation_rate']) * 100:.2f} 个百分点，代价是放弃一部分尚未得到多次验证的新区域。")
    _add_figure(document, result_dir / "semantic_update_safety.png", "图 1  四种语义更新策略的可接受覆盖率与 Oct.12 未来细类别确认率")

    _add_heading(document, "四、可落地的地图更新规则")
    _add_body(
        document,
        "因此，项目中的语义地图更新应采用延迟确认而不是即时写入：当前帧仅为单元积累候选证据；当该单元当前为静态类别、历史稳定度和证据数量达标、当前细类别与历史主导细类别一致，并在跨帧重复观察后，才提高该单元的语义可靠性权重。对人员、叉车、推车、货物等动态或半动态类别，只记录短期观测，不写入长期层；对暂未被历史地图覆盖的新单元，可建立待确认缓冲区，达到同样的跨帧和跨次观测条件后再转为正式稳定单元。")

    _add_heading(document, "五、结论")
    _add_body(
        document,
        "本实验完成了长期更新机制的离线安全性验证：严格门控能以可量化的覆盖率代价换取更高的未来一致性，且整个判断过程没有读取未来日期。它不直接声称会提升某一次定位误差，而是为后续长周期运行提供了“不让临时物体污染地图”的准入规则。结合实验四的多帧静态配准，系统可以同时做到短期内利用当前与过去帧恢复被遮挡的几何证据，长期内只把重复证实的静态语义吸收到可靠性层中。")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the Experiment 5 research log from the July 18 template.")
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--result-dir", required=True)
    args = parser.parse_args()
    build_log(Path(args.template), Path(args.output), Path(args.result_dir))


if __name__ == "__main__":
    main()
